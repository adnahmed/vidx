"""Excel scene import: parse, validate and normalize rows.

Excel is import-only. Once rows are persisted as scenes, the database is the
single source of truth and the workbook is never consulted again.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, List, Optional

from openpyxl import load_workbook


@dataclass
class ParsedSceneRow:
    """One validated scene row from the workbook."""

    scene_number: int
    image_prompt: str = ""
    scene_prompt: str = ""
    audio_prompt: str = ""


@dataclass
class ImportResult:
    """Outcome of parsing an uploaded workbook."""

    rows: List[ParsedSceneRow] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    sheet_name: str = ""


IMAGE_PROMPT_HEADERS = {
    "image prompt",
    "imageprompt",
    "image_prompt",
    "image",
}
SCENE_PROMPT_HEADERS = {
    "scene prompt",
    "sceneprompt",
    "scene_prompt",
    "scene",
    "video prompt",
    "video_prompt",
    "videoprompt",
}
AUDIO_PROMPT_HEADERS = {
    "audio prompt",
    "audioprompt",
    "audio_prompt",
    "audio",
    "voice prompt",
    "voice_prompt",
    "voiceprompt",
}
SCENE_NUMBER_HEADERS = {
    "scene number",
    "scenenumber",
    "scene_number",
    "scene #",
    "scene no",
    "scene no.",
    "#",
    "no",
    "no.",
    "number",
}


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    return " ".join(_normalize(value).lower().replace("-", " ").split())


def _coerce_scene_number(value: Any) -> Optional[int]:
    text = _normalize(value)
    if not text:
        return None
    try:
        number = int(float(text))
    except (TypeError, ValueError):
        return None
    return number


def parse_scene_workbook(data: bytes) -> ImportResult:
    """Parse an uploaded .xlsx workbook into validated scene rows."""
    result = ImportResult()
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        result.errors.append(f"Could not read the Excel file: {exc}")
        return result

    sheet = workbook[workbook.sheetnames[0]]
    result.sheet_name = sheet.title

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        result.errors.append("The Excel file is empty.")
        return result

    header_index = None
    headers: List[str] = []
    for index, row in enumerate(rows):
        if any(_normalize(cell) for cell in row):
            header_index = index
            headers = [_normalize_header(cell) for cell in row]
            break
    if header_index is None:
        result.errors.append("The Excel file is empty.")
        return result

    column_map: dict[str, int] = {}
    for position, header in enumerate(headers):
        if header in IMAGE_PROMPT_HEADERS and "image_prompt" not in column_map:
            column_map["image_prompt"] = position
        elif header in SCENE_NUMBER_HEADERS and "scene_number" not in column_map:
            column_map["scene_number"] = position
        elif header in SCENE_PROMPT_HEADERS and "scene_prompt" not in column_map:
            column_map["scene_prompt"] = position
        elif header in AUDIO_PROMPT_HEADERS and "audio_prompt" not in column_map:
            column_map["audio_prompt"] = position

    if not {"image_prompt", "scene_prompt", "audio_prompt"} & set(column_map):
        result.errors.append(
            "No recognizable prompt columns found. Expected columns named "
            "'Image Prompt', 'Scene Prompt' and/or 'Audio Prompt' "
            "(optionally 'Scene Number').",
        )
        return result

    seen_numbers: set[int] = set()
    derived_number = 0
    for row_index in range(header_index + 1, len(rows)):
        row = rows[row_index]
        if not any(_normalize(cell) for cell in row):
            continue

        def cell(key: str, row_values: tuple = row) -> str:
            position = column_map.get(key)
            if position is None or position >= len(row_values):
                return ""
            return _normalize(row_values[position])

        image_prompt = cell("image_prompt")
        scene_prompt = cell("scene_prompt")
        audio_prompt = cell("audio_prompt")
        raw_number = cell("scene_number")
        excel_row = row_index + 1

        if not (image_prompt or scene_prompt or audio_prompt):
            result.errors.append(
                f"Row {excel_row}: at least one prompt (image, scene or audio) is required.",
            )
            continue

        if raw_number:
            number = _coerce_scene_number(raw_number)
            if number is None or number <= 0:
                result.errors.append(
                    f"Row {excel_row}: scene number '{raw_number}' is not a positive integer.",
                )
                continue
        else:
            derived_number += 1
            number = derived_number

        if number in seen_numbers:
            result.errors.append(f"Row {excel_row}: duplicate scene number {number}.")
            continue
        seen_numbers.add(number)

        result.rows.append(
            ParsedSceneRow(
                scene_number=number,
                image_prompt=image_prompt,
                scene_prompt=scene_prompt,
                audio_prompt=audio_prompt,
            ),
        )

    if not result.rows and not result.errors:
        result.errors.append("No scene rows found in the workbook.")
    return result
